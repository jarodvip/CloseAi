import sys
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

from scripts.import_docs import parse_any, read_docx  # noqa: E402


def _make_docx(tmp_path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("第一段：气泡水行业景气度高。" )
    doc.add_paragraph("第二段：" + "渠道下沉" * 150)
    p = tmp_path / "sample.docx"
    doc.save(str(p))
    return p


def test_read_docx(tmp_path):
    p = _make_docx(tmp_path)
    text = read_docx(p)
    assert "气泡水行业景气度高" in text
    assert "渠道下沉" in text


def test_parse_any_dispatch_and_bad_file(tmp_path):
    p = _make_docx(tmp_path)
    assert "气泡水" in parse_any(p)
    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"not-a-real-docx")
    assert parse_any(bad) == ""  # 解析失败返回空串而非抛出


def test_parse_unsupported_ext(tmp_path):
    f = tmp_path / "x.exe"
    f.write_bytes(b"\x00")
    assert parse_any(f) == ""


# CLI 混批回归用的唯一标记：入库行全部挂在该 source_name 下，finally 删尽
MARKER_SOURCE = "import_docs_zz9"
MARKER_INDUSTRY = "导入CLI"


def _make_docx_at(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("第一段：气泡水行业景气度高。")
    doc.add_paragraph("第二段：" + "渠道下沉" * 50)
    doc.save(str(path))
    return path


def _purge_marker_rows():
    """删尽本测试在真实 dev DB 中创建的全部行（含失败路径），并自证清零"""
    from sqlalchemy import text

    from app.core.domain.services.research_service import delete_chunk
    from app.db.session import SessionLocal

    db = SessionLocal()
    try:
        ids = db.execute(
            text("SELECT id FROM research_chunks WHERE source_name = :s"),
            {"s": MARKER_SOURCE},
        ).fetchall()
        for (cid,) in ids:
            assert delete_chunk(db, cid)
        left = db.execute(
            text("SELECT COUNT(*) FROM research_chunks WHERE source_name = :s"),
            {"s": MARKER_SOURCE},
        ).scalar()
        assert left == 0
    finally:
        db.close()


def test_main_batch_survives_single_file_ingest_failure(tmp_path, monkeypatch):
    """CLI 混批回归：单个文件入库抛错须告警回滚并继续，绝不让整批崩掉"""
    import scripts.import_docs as mod

    _make_docx_at(tmp_path / "flaky.docx")   # 入库阶段抛故障
    _make_docx_at(tmp_path / "good.docx")    # 必须照常入库成功
    (tmp_path / "broken.pdf").write_bytes(b"garbage-not-a-pdf")  # 解析阶段即挂

    real_ingest = mod.ingest_text
    attempted = []

    def flaky_ingest(db, content, **kwargs):
        attempted.append(kwargs.get("title"))
        if kwargs.get("title") == "flaky":
            raise RuntimeError("模拟入库故障")
        return real_ingest(db, content, **kwargs)

    monkeypatch.setattr(mod, "ingest_text", flaky_ingest)

    try:
        # 单文件入库失败不得向外抛异常（旧行为在此处炸出 RuntimeError）
        mod.main([str(tmp_path), "--industry", MARKER_INDUSTRY,
                  "--source-name", MARKER_SOURCE])
        assert "flaky" in attempted   # 故障文件确实被尝试过

        from sqlalchemy import text

        from app.db.session import SessionLocal

        db = SessionLocal()
        try:
            rows = db.execute(
                text("SELECT title FROM research_chunks WHERE source_name = :s"),
                {"s": MARKER_SOURCE},
            ).fetchall()
        finally:
            db.close()
        titles = {t for (t,) in rows}
        assert "good" in titles       # 好文件仍完整入库，循环未被中断
        assert "flaky" not in titles  # 故障文件未留下半截行
    finally:
        _purge_marker_rows()
